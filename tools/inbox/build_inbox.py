#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""构建「分享收件箱」：合并 index.jsonl + 手工条目 + 待处理口令，产出 CSV / XLSX / HTML 看板。

数据源：
  - $PUBLIC_LINK_KNOWLEDGE_DIR/index.jsonl   （已抓取知识包）
  - ~/.public-link-knowledge/inbox/inbox_extra.jsonl  （手工补录：未建包但已解读的分享/tips）
  - ~/.public-link-knowledge/pending-reviews.md        （仅口令、缺链接的待办）
输出（~/.public-link-knowledge/inbox/）：
  - inbox.csv / inbox.xlsx / inbox.html（本地看板）
用法：python3 build_inbox.py
"""
import csv
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from inbox_common import norm_url, write_lock  # noqa: E402
from inbox_state import STATE_CN, fold, read_events  # noqa: E402

INBOX = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(INBOX)
# 知识库目录：环境变量优先；默认 ROOT/public-links
KNOWLEDGE = os.environ.get("PUBLIC_LINK_KNOWLEDGE_DIR") or os.path.join(ROOT, "public-links")
INDEX = os.path.join(KNOWLEDGE, "index.jsonl")
EXTRA = os.path.join(INBOX, "inbox_extra.jsonl")
PENDING = os.path.join(ROOT, "pending-reviews.md")

SOURCE_CN = {"douyin": "抖音", "wechat": "公众号", "web": "网页", "x": "X",
             "xiaohongshu": "小红书", "manual": "手工", "weixin": "微信",
             "bilibili": "B站", "reddit": "Reddit", "youtube": "YouTube",
             "github": "GitHub", "hackernews": "HN", "wechat_video": "视频号", "asr": "ASR"}
AGENT_CN = {"hermes": "Hermes", "cc": "Claude Code", "claude": "Claude Code", "claude-code": "Claude Code",
            "codex": "Codex", "agy": "agy", "antigravity": "agy", "grok": "Grok",
            "grokbot": "Grok Bot", "grok-bot": "Grok Bot", "muse": "Muse", "user": "用户"}
CHANNEL_CN = {"wechat": "微信", "weixin": "微信", "manual": "手动", "test": "测试", "cron": "定时"}


def label(value, table):
    """来源 agent / 渠道的显示名；没记录的统一显示为空（看板上归入"未知"）。"""
    v = (value or "").strip()
    return table.get(v.lower(), v)


STATUS_CN = {"extracted": "已抓取", "partial": "部分抓取", "failed": "抓取失败",
             "needs_browser": "需浏览器", "no_link": "待补链接", "tip": "提示词/要点", "manual": "手工记录"}


def classify(rec):
    """粗略推断条目类型。"""
    t = (rec.get("title") or "") + " " + (rec.get("summary") or "")
    if any(k in t for k in ("提示词", "prompt", "提示语")):
        return "提示词"
    if any(k in t for k in ("skill", "Skill", "技能")):
        return "Skill"
    if any(k in t for k in ("工具", "软件", "开源", "GitHub", "github")):
        return "工具/开源"
    if rec.get("platform") in ("douyin", "xiaohongshu", "weixin"):
        return "视频/图文"
    return "文章"


def load_index():
    rows = []
    if not os.path.exists(INDEX):
        return rows
    for line_no, line in enumerate(open(INDEX, encoding="utf-8"), 1):
        line = line.strip()
        if not line:
            continue
        try:
            d = json.loads(line)
        except json.JSONDecodeError:
            print(f"警告: 知识库 index.jsonl 第{line_no}行不是合法 JSON，已跳过")
            continue
        if not isinstance(d, dict):
            print(f"警告: 知识库 index.jsonl 第{line_no}行不是 JSON 对象，已跳过")
            continue
        rows.append({
            # 日期取首次入库时间；重抓/重试只刷新 captured_at，不该让旧链接冒充"今天"
            "date": (d.get("first_seen") or d.get("captured_at") or "")[:10],
            "agent": label(d.get("agent"), AGENT_CN),
            "channel": label(d.get("channel"), CHANNEL_CN),
            "source": SOURCE_CN.get(d.get("platform"), d.get("platform") or ""),
            "title": (d.get("title") or "").strip(),
            "author": (d.get("author") or "").strip(),
            "url": d.get("url") or d.get("resolved_url") or "",
            "type": classify(d),
            "stats": "",
            "summary": "",
            "status": STATUS_CN.get(d.get("capture_status"), d.get("capture_status") or ""),
            "tags": "",
            "package": d.get("package") or "",
            "notes": "",
            "detail": "",
        })
    return rows


def load_extra():
    rows = []
    if not os.path.exists(EXTRA):
        return rows
    for line_no, line in enumerate(open(EXTRA, encoding="utf-8"), 1):
        line = line.strip()
        if not line:
            continue
        try:
            d = json.loads(line)
        except json.JSONDecodeError:
            print(f"警告: inbox_extra.jsonl 第{line_no}行不是合法 JSON，已跳过")
            continue
        if not isinstance(d, dict):
            print(f"警告: inbox_extra.jsonl 第{line_no}行不是 JSON 对象，已跳过")
            continue
        if isinstance(d.get("tags"), list):
            d["tags"] = ",".join(str(t) for t in d["tags"])
        d["agent"] = label(d.get("agent"), AGENT_CN)
        d["channel"] = label(d.get("channel"), CHANNEL_CN)
        rows.append(d)
    return rows


def load_pending():
    """解析 pending-reviews.md「未处理」里的口令条目。"""
    rows = []
    if not os.path.exists(PENDING):
        return rows
    txt = open(PENDING, encoding="utf-8").read()
    block = txt.split("## 未处理", 1)[1].split("## 已处理", 1)[0] if "## 未处理" in txt else ""
    # 只认「- [ ] **标题**」形式的待办条目；账本里以 ">" 开头的每夜汇总笔记里的加粗片段不算条目
    for line in block.splitlines():
        line = line.strip()
        if not re.match(r"-\s*\[[ x]\]", line):
            continue
        title = ""
        m = re.search(r"\*\*(.+?)\*\*", line)
        if m:
            title = m.group(1).strip()
        if title:
            rows.append({
                "date": "", "source": "抖音", "title": title, "author": "",
                "url": "", "type": "待回填", "stats": "", "summary": "仅分享口令，缺 v.douyin.com 短链，无法解析",
                "status": "待补链接", "tags": "待补链", "package": "", "notes": "", "detail": "",
            })
    return rows


def dedupe(rows):
    """按 URL 去重；重复时保留信息更全的一条（extra 优先补中文标题/摘要）。"""
    order, bykey = [], {}
    for r in rows:
        key = norm_url(r.get("url")) or ("T:" + (r.get("title") or "")[:40])
        if key not in bykey:
            bykey[key] = dict(r)
            order.append(key)
        else:
            cur = bykey[key]
            # 同一 URL 多次抓取：状态更好的（或同级更新的）覆盖旧状态，避免"后来抓成功了仍显示失败"
            if STATUS_RANK.get(r.get("status"), 1) >= STATUS_RANK.get(cur.get("status"), 1) \
                    and r.get("source") == cur.get("source"):
                for f in ("status", "package"):
                    if r.get(f):
                        cur[f] = r[f]
            for f in FIELDS:
                if not cur.get(f) and r.get(f):
                    cur[f] = r[f]
            if (cur.get("title", "").startswith("http")
                    and r.get("title") and not r["title"].startswith("http")):
                cur["title"] = r["title"]
    out = [bykey[k] for k in order]
    out.sort(key=lambda r: (r.get("date") or "0000-00-00"), reverse=True)
    return out


STATUS_RANK = {"已抓取": 3, "提示词/要点": 3, "手工记录": 2, "部分抓取": 1,
               "需浏览器": 0, "待补链接": 0, "抓取失败": -1}
FIELDS = ["date", "source", "title", "author", "url", "type", "stats",
          "summary", "status", "tags", "package", "notes", "detail", "agent", "channel",
          "state", "due", "owner"]
HDR = {"state": "进度", "due": "到期", "owner": "负责人", "agent": "来源Agent", "channel": "渠道","date": "日期", "source": "来源", "title": "标题", "author": "作者",
       "url": "链接", "type": "类型", "stats": "热度(赞/藏/评)", "summary": "摘要/要点",
       "status": "状态", "tags": "标签", "package": "知识包", "notes": "备注", "detail": "详细记录"}


def atomic_write_text(path, text, encoding="utf-8"):
    """先写临时文件再改名：两个 agent 同时重建看板时，不会出现一份文件被另一份的尾部污染。"""
    tmp = "%s.%d.tmp" % (path, os.getpid())
    with open(tmp, "w", encoding=encoding) as f:
        f.write(text)
    os.replace(tmp, path)


def write_csv(rows, path):
    tmp = "%s.%d.tmp" % (path, os.getpid())
    with open(tmp, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([HDR[k] for k in FIELDS])
        for r in rows:
            w.writerow([r.get(k, "") for k in FIELDS])
    os.replace(tmp, path)


def write_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "分享收件箱"
    ws.append([HDR[k] for k in FIELDS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B6EA5")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([r.get(k, "") for k in FIELDS])
    widths = {"date": 12, "source": 8, "title": 56, "author": 16, "url": 42,
              "type": 12, "stats": 18, "summary": 60, "status": 12, "tags": 16,
              "package": 30, "notes": 20, "detail": 42, "agent": 12, "channel": 8,
              "state": 10, "due": 12, "owner": 12}
    for i, k in enumerate(FIELDS, 1):
        ws.column_dimensions[chr(64 + i)].width = widths[k]
    ws.freeze_panes = "A2"
    tmp = "%s.%d.tmp.xlsx" % (path, os.getpid())
    wb.save(tmp)
    os.replace(tmp, path)
    return True


def write_html(rows, path):
    """用模板渲染看板（模板见 inbox_template.html，AIHOT 风格）。"""
    tpl_path = os.path.join(INBOX, "inbox_template.html")
    tpl = open(tpl_path, encoding="utf-8").read()
    # 数据嵌在 <script> 里：标题等来自抓取内容，必须转义 < > &，否则 "</script>" 能跳出脚本执行任意代码
    def safe(obj):
        return json.dumps(obj, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    html = tpl.replace("__DATA__", safe(rows)).replace("__HDR__", safe([HDR[k] for k in FIELDS]))
    atomic_write_text(path, html)


def attach_state(rows):
    """把 inbox_state.jsonl 折叠出的进度（待分拣/处理中/等待中…）、到期日、负责人并入每一行。"""
    folded = fold(read_events())
    for r in rows:
        key = norm_url(r.get("url")) or ("T:" + (r.get("title") or "")[:40])
        rec = folded.get(key)
        state = rec["state"] if rec else "inbox"
        r["state"] = STATE_CN.get(state, state)
        r["due"] = rec["due"] if rec else ""
        r["owner"] = label(rec["owner"], AGENT_CN) if rec else ""
    return rows


def main():
    os.makedirs(INBOX, exist_ok=True)
    with write_lock(".build.lock"):
        build()


def build():
    rows = attach_state(dedupe(load_extra() + load_index() + load_pending()))
    write_csv(rows, os.path.join(INBOX, "inbox.csv"))
    ok = write_xlsx(rows, os.path.join(INBOX, "inbox.xlsx"))
    write_html(rows, os.path.join(INBOX, "inbox.html"))
    by = {}
    for r in rows:
        name = r.get("source") or "网页"
        by[name] = by.get(name, 0) + 1
    print(f"条目总数: {len(rows)}")
    print("按来源:", ", ".join(f"{k} {v}" for k, v in sorted(by.items(), key=lambda x: -x[1])))
    bad = [r for r in rows if r.get("status") in ("抓取失败", "需浏览器")]
    if bad:
        print(f"待重试: {len(bad)} 条（python3 inbox_retry.py 查看 / --run 重抓）")
    print("xlsx:", "已生成" if ok else "跳过(缺 openpyxl)")
    print("输出目录:", INBOX)


if __name__ == "__main__":
    main()
